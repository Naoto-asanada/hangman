"""DADS（デジタル庁デザインシステム）準拠のスライド部品。

座標・サイズはすべて px（スライド = 1280×720px、1px = 1/96in）で指定する。
文字サイズも px で指定し、内部で pt（px × 0.75）に変換する。
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from lxml import etree

# ---- トークン（references/tokens.md と同期させること） -------------------
FONT = "Noto Sans JP"

TEXT = "1A1A1A"
TEXT_SUB = "595959"
PRIMARY = "0017C1"
PRIMARY_BG = "E8F1FE"
LINE = "949494"
LINE_STRONG = "1A1A1A"
LINE_ON_TINT = "767676"
DANGER = "CE0000"
SUCCESS = "197A4B"
BG = "FFFFFF"

SIZE_TITLE = 32
SIZE_H2 = 22
SIZE_H3 = 18
SIZE_BODY = 16
SIZE_NOTE = 14

W, H = 1280, 720
MX = 64            # 左右の余白
TOP = 48           # タイトル上端
BODY_Y = 144       # 本文開始
BODY_BOTTOM = 672
GUTTER = 32
CONTENT_W = W - MX * 2  # 1152


def px(v):
    return Emu(int(round(v * 9525)))


def cols(n, x=MX, width=CONTENT_W, gutter=GUTTER):
    """n 列の (x, w) を 8px 単位で返す。"""
    w = (width - gutter * (n - 1)) // n // 8 * 8
    return [(x + i * (w + gutter), w) for i in range(n)]


# ---- 基本部品 ---------------------------------------------------------------
def new_deck():
    prs = Presentation()
    prs.slide_width = px(W)
    prs.slide_height = px(H)
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _set_font(run, size, bold=False, color=TEXT, underline=False):
    f = run.font
    f.size = Pt(size * 0.75)
    f.bold = bold
    f.underline = underline
    f.name = FONT
    f.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def text(slide, x, y, w, h, content, size=SIZE_BODY, bold=False, color=TEXT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.6,
         bullets=False, para_gap=None):
    """テキストボックス。content は文字列か、文字列のリスト（段落）。

    段落の要素を (文字列, dict) にすると段落ごとに size/bold/color を上書きできる。
    """
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = content if isinstance(content, list) else [content]
    for i, p_ in enumerate(paras):
        opts = {}
        if isinstance(p_, tuple):
            p_, opts = p_
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = max(line_spacing, 1.5)
        if i > 0 and para_gap is not None:
            p.space_before = Pt(para_gap * 0.75)
        r = p.add_run()
        r.text = ("・" + p_) if bullets else p_
        _set_font(r, opts.get("size", size), opts.get("bold", bold),
                  opts.get("color", color))
    return tb


def _drop_style(shape):
    """テーマ由来の影・効果を残さないよう p:style を外す（塗り・線は明示指定する）。"""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def rect(slide, x, y, w, h, fill=None, line=None, line_px=1, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        px(x), px(y), px(w), px(h))
    if radius:
        shape.adjustments[0] = min(8 / min(w, h), 0.5)  # 角丸 8px
    shape.shadow.inherit = False
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = px(line_px)
    else:
        shape.line.fill.background()
    shape.text_frame.text = ""
    _drop_style(shape)
    return shape


def hline(slide, x, y, w, color=LINE, width=1):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x), px(y), px(x + w), px(y))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = px(width)
    _drop_style(c)
    return c


def vline(slide, x, y, h, color=LINE, width=1):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x), px(y), px(x), px(y + h))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = px(width)
    _drop_style(c)
    return c


def label(slide, x, y, content, kind="primary"):
    """文字ラベル（色だけに頼らないための必須部品）。幅は文字数から推定。"""
    color = {"primary": PRIMARY, "danger": DANGER, "success": SUCCESS,
             "neutral": TEXT_SUB}[kind]
    w = (len(content) * SIZE_BODY + 24) // 8 * 8 + 8
    rect(slide, x, y, w, 32, fill=color, radius=True)
    text(slide, x, y, w, 32, content, size=SIZE_BODY, bold=True, color=BG,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.5)
    return w


def slide_title(slide, title, lead=None):
    text(slide, MX, TOP, CONTENT_W, 48, title, size=SIZE_TITLE, bold=True, line_spacing=1.5)
    if lead:
        text(slide, MX, TOP + 56, CONTENT_W, 32, lead, size=SIZE_BODY, color=TEXT_SUB)


def note(slide, content):
    """出典・注記（14px はここだけで使う）。"""
    text(slide, MX, BODY_BOTTOM + 8, CONTENT_W, 24, content, size=SIZE_NOTE,
         color=TEXT_SUB, line_spacing=1.5)


# ---- 表 ---------------------------------------------------------------------
def _cell_border(cell, side, width_px, color):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"L": "a:lnL", "R": "a:lnR", "T": "a:lnT", "B": "a:lnB"}[side]
    old = tcPr.find(qn(tag))
    if old is not None:
        tcPr.remove(old)
    ln = etree.Element(qn(tag), w=str(int(width_px * 9525)), cap="flat", cmpd="sng", algn="ctr")
    if width_px:
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=color)
    else:
        etree.SubElement(ln, qn("a:noFill"))
    # スキーマ順 lnL, lnR, lnT, lnB を守って先頭側に挿入
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    idx = 0
    for t in order[:order.index(tag)]:
        if tcPr.find(qn(t)) is not None:
            idx += 1
    tcPr.insert(idx, ln)


def _is_number(s):
    t = s.replace(",", "").replace("%", "").replace("pt", "").replace("円", "")
    t = t.replace("＋", "").replace("−", "").replace("+", "").replace("-", "").replace("億", "").replace("万", "")
    try:
        float(t)
        return True
    except ValueError:
        return False


def table(slide, x, y, w, rows, col_widths=None, highlight_col=None,
          highlight_row=None, highlight_label="推奨", row_h=48):
    """DADS の表：見出し行は太字＋黒線、セル上寄せ、数値右寄せ、推奨だけ薄青。"""
    n_rows, n_cols = len(rows), len(rows[0])
    gf = slide.shapes.add_table(n_rows, n_cols, px(x), px(y), px(w), px(row_h * n_rows))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.set(attr, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        tblPr.remove(style)
    widths = col_widths or [w // n_cols] * n_cols
    for i, cw in enumerate(widths):
        tbl.columns[i].width = px(cw)
    for r in range(n_rows):
        tbl.rows[r].height = px(row_h)
        for c in range(n_cols):
            cell = tbl.cell(r, c)
            val = str(rows[r][c])
            hl = (c == highlight_col and highlight_col is not None) or \
                 (r == highlight_row and highlight_row is not None)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(PRIMARY_BG if hl else BG)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            cell.margin_left = cell.margin_right = px(16)
            cell.margin_top = cell.margin_bottom = px(8)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.5
            is_head = r == 0
            if highlight_label in val:
                pass
            elif is_head and hl and highlight_col == c:
                val = f"{val}【{highlight_label}】"
            if not is_head and c == 0 and hl and highlight_row == r:
                val = f"{val}【{highlight_label}】"
            p.alignment = PP_ALIGN.RIGHT if (r > 0 and c > 0 and _is_number(val)) else PP_ALIGN.LEFT
            run = p.add_run()
            run.text = val
            _set_font(run, SIZE_BODY, bold=is_head or (c == 0),
                      color=PRIMARY if (is_head and hl) else TEXT)
            for side in ("L", "R", "T"):
                _cell_border(cell, side, 0, LINE)
            _cell_border(cell, "B", 2 if is_head else 1, LINE_STRONG if is_head else LINE)
    return gf


# ---- 型 ---------------------------------------------------------------------
def title_slide(prs, title, subtitle="", meta=""):
    s = blank(prs)
    text(s, MX, 248, CONTENT_W, 96, title, size=40, bold=True, line_spacing=1.4)
    if subtitle:
        text(s, MX, 360, CONTENT_W, 32, subtitle, size=SIZE_H3, color=TEXT_SUB)
    hline(s, MX, 424, CONTENT_W)
    if meta:
        text(s, MX, 448, CONTENT_W, 32, meta, size=SIZE_BODY, color=TEXT_SUB)
    return s


def section_slide(prs, number, name):
    s = blank(prs)
    text(s, MX, 280, CONTENT_W, 56, f"{number:02d}", size=40, bold=True, color=PRIMARY)
    text(s, MX, 344, CONTENT_W, 56, name, size=SIZE_TITLE, bold=True)
    return s


def index_slide(prs, title, entries):
    """entries: [(項目名, 遷移先スライド)]。青文字＋下線のリンクで並べる。
    遷移先スライドが後で作られる場合は、作成後に link_index() で接続する。"""
    s = blank(prs)
    slide_title(s, title)
    boxes = []
    per_col = 10
    for i, (name, _) in enumerate(entries):
        col, row = divmod(i, per_col)
        (cx, cw) = cols(2)[col]
        tb = s.shapes.add_textbox(px(cx), px(BODY_Y + row * 48), px(cw), px(32))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.5
        r = p.add_run()
        r.text = f"{i + 1:02d}  {name}"
        _set_font(r, SIZE_BODY, color=PRIMARY, underline=True)
        boxes.append(tb)
    s._dads_index = boxes
    link_index(s, [t for _, t in entries])
    return s


def link_index(index_slide_, targets):
    for tb, target in zip(index_slide_._dads_index, targets):
        if target is not None:
            tb.click_action.target_slide = target


def exec_summary_slide(prs, title, points, metrics, lead=None):
    """左に要点（最大4）、右に主要数値（最大3: (数値, 説明)）。"""
    s = blank(prs)
    slide_title(s, title, lead)
    left_w, right_x = 712, MX + 712 + GUTTER
    right_w = W - MX - right_x
    y = BODY_Y + 16
    for i, pt in enumerate(points[:4]):
        text(s, MX, y, 40, 32, f"{i + 1}", size=SIZE_H3, bold=True, color=PRIMARY)
        text(s, MX + 40, y, left_w - 40, 104, pt, size=SIZE_BODY)
        y += 120
    vline(s, right_x - GUTTER // 2, BODY_Y + 16, 480)
    my = BODY_Y + 16
    for i, (num, desc) in enumerate(metrics[:3]):
        if i:
            hline(s, right_x + 16, my - 16, right_w - 16)
        text(s, right_x + 16, my, right_w - 16, 72, num, size=48, bold=True, color=PRIMARY, line_spacing=1.5)
        text(s, right_x + 16, my + 80, right_w - 16, 48, desc, size=SIZE_BODY, color=TEXT_SUB)
        my += 168
    return s


def conclusion_reasons_slide(prs, title, conclusion, reasons):
    """上に結論、線、下に根拠3つ（(見出し, 本文)）。枠は使わない。"""
    s = blank(prs)
    slide_title(s, title)
    text(s, MX, BODY_Y, CONTENT_W, 96, conclusion, size=SIZE_H2, bold=True)
    hline(s, MX, BODY_Y + 120, CONTENT_W)
    for i, ((cx, cw), (head, body)) in enumerate(zip(cols(len(reasons)), reasons)):
        y = BODY_Y + 152
        text(s, cx, y, cw, 32, f"根拠{i + 1}", size=SIZE_BODY, bold=True, color=PRIMARY)
        text(s, cx, y + 32, cw, 64, head, size=SIZE_H3, bold=True)
        text(s, cx, y + 104, cw, 240, body, size=SIZE_BODY)
    return s


def kpi_slide(prs, title, kpis, lead=None, body=None):
    """kpis: [(名前, 値, 増減テキスト, 状態)] 状態は 'good' / 'bad' / 'flat'。
    増減は必ず文字（＋8% など）。悪化だけ赤ラベル「悪化」。"""
    s = blank(prs)
    slide_title(s, title, lead)
    for (cx, cw), (name, value, delta, state) in zip(cols(len(kpis)), kpis):
        y = BODY_Y + 8
        rect(s, cx, y, cw, 192, line=LINE, radius=True)
        text(s, cx + 24, y + 24, cw - 48, 32, name, size=SIZE_BODY, color=TEXT_SUB)
        text(s, cx + 24, y + 56, cw - 48, 64, value, size=40, bold=True, line_spacing=1.5)
        text(s, cx + 24, y + 136, cw - 48, 32, f"前期比 {delta}", size=SIZE_BODY,
             bold=True, color=DANGER if state == "bad" else TEXT)
        if state == "bad":
            label(s, cx + cw - 24 - 64, y + 24, "悪化", "danger")
    if body:
        text(s, MX, BODY_Y + 232, CONTENT_W, 280, body, size=SIZE_BODY, bullets=True, para_gap=8)
    return s


def table_slide(prs, title, rows, lead=None, source=None, **kw):
    s = blank(prs)
    slide_title(s, title, lead)
    table(s, MX, BODY_Y + 8, CONTENT_W, rows, **kw)
    if source:
        note(s, source)
    return s


def cards_slide(prs, title, cards, recommended=None, lead=None):
    """cards: [(名前, 主要値, [特徴...])]。輪郭線あり・影なし。推奨だけ薄青＋ラベル。"""
    s = blank(prs)
    slide_title(s, title, lead)
    for i, ((cx, cw), (name, value, feats)) in enumerate(zip(cols(len(cards)), cards)):
        rec = i == recommended
        y = BODY_Y + 8
        rect(s, cx, y, cw, 456, fill=PRIMARY_BG if rec else None,
             line=LINE_ON_TINT if rec else LINE, radius=True)
        if rec:
            label(s, cx + 24, y + 24, "推奨")
        text(s, cx + 24, y + 56, cw - 48, 40, name, size=SIZE_H2, bold=True)
        text(s, cx + 24, y + 104, cw - 48, 56, value, size=32, bold=True, color=PRIMARY, line_spacing=1.5)
        hline(s, cx + 24, y + 176, cw - 48, color=LINE_ON_TINT if rec else LINE)
        text(s, cx + 24, y + 200, cw - 48, 232, feats, size=SIZE_BODY, bullets=True, para_gap=8)
    return s


def steps_slide(prs, title, steps, lead=None):
    """steps: [(見出し, 本文)] 最大5。番号付き横並び、細い灰色線でつなぐ。"""
    s = blank(prs)
    slide_title(s, title, lead)
    cs = cols(len(steps))
    hline(s, MX + 24, BODY_Y + 40, CONTENT_W - 48)
    for i, ((cx, cw), (head, body)) in enumerate(zip(cs, steps)):
        rect(s, cx, BODY_Y + 16, 48, 48, fill=PRIMARY, radius=True)
        text(s, cx, BODY_Y + 16, 48, 48, f"{i + 1}", size=SIZE_H3, bold=True, color=BG,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, cx, BODY_Y + 88, cw, 64, head, size=SIZE_H3, bold=True)
        text(s, cx, BODY_Y + 160, cw, 320, body, size=SIZE_BODY)
    return s


def flow3_slide(prs, title, inputs, core, outputs, lead=None):
    """入力／処理／出力の3列。中心の処理基盤だけ薄青面。"""
    s = blank(prs)
    slide_title(s, title, lead)
    (x1, w1), (x2, w2), (x3, w3) = cols(3)
    heads = [("入力", x1, w1), ("処理", x2, w2), ("出力", x3, w3)]
    for h_, x, w in heads:
        text(s, x, BODY_Y, w, 32, h_, size=SIZE_H3, bold=True)
        hline(s, x, BODY_Y + 40, w, color=LINE_STRONG, width=2)

    def items(x, w, lst):
        y = BODY_Y + 64
        for it in lst:
            rect(s, x, y, w, 64, line=LINE, radius=True)
            text(s, x + 16, y, w - 32, 64, it, size=SIZE_BODY, anchor=MSO_ANCHOR.MIDDLE)
            y += 80
    items(x1, w1, inputs)
    items(x3, w3, outputs)
    name, details = core
    h_core = max(80 * max(len(inputs), len(outputs)) - 16, 224)
    rect(s, x2, BODY_Y + 64, w2, h_core, fill=PRIMARY_BG, line=LINE_ON_TINT, radius=True)
    text(s, x2 + 24, BODY_Y + 88, w2 - 48, 40, name, size=SIZE_H3, bold=True, color=PRIMARY)
    text(s, x2 + 24, BODY_Y + 136, w2 - 48, h_core - 96, details, size=SIZE_BODY, bullets=True)
    return s


def tree_slide(prs, title, goal, drivers, lead=None):
    """goal: 文字列。drivers: [(指標, 重点?, [(施策, 重点?)])]。重点だけ薄青＋ラベル。"""
    s = blank(prs)
    slide_title(s, title, lead)
    gx, gw = MX, 240
    mx_, mw = MX + 240 + 64, 352
    ax, aw = mx_ + mw + 64, W - MX - (mx_ + mw + 64)
    n_leaves = sum(max(len(a), 1) for _, _, a in drivers)
    leaf_h, leaf_gap = 48, 8
    total = n_leaves * leaf_h + (n_leaves - 1) * leaf_gap
    top = BODY_Y + 8
    gy = (top + total // 2 - 40) // 8 * 8
    rect(s, gx, gy, gw, 80, line=LINE_STRONG, radius=True)
    text(s, gx + 16, gy, gw - 32, 80, goal, size=SIZE_H3, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    y = top
    for metric, m_focus, actions in drivers:
        n = max(len(actions), 1)
        block_h = n * leaf_h + (n - 1) * leaf_gap
        my = (y + block_h // 2 - 24) // 8 * 8
        hline(s, gx + gw, gy + 40, 32)
        vline(s, gx + gw + 32, min(gy + 40, my + 24), abs(gy + 40 - (my + 24)))
        hline(s, gx + gw + 32, my + 24, 32)
        rect(s, mx_, my, mw, 48, fill=PRIMARY_BG if m_focus else None,
             line=LINE_ON_TINT if m_focus else LINE, radius=True)
        text(s, mx_ + 16, my, mw - 96, 48, metric, size=SIZE_BODY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        if m_focus:
            label(s, mx_ + mw - 80, my + 8, "重点")
        for j, (act, a_focus) in enumerate(actions):
            ay = y + j * (leaf_h + leaf_gap)
            hline(s, mx_ + mw, my + 24, 32)
            vline(s, mx_ + mw + 32, min(my + 24, ay + 24), abs(my + 24 - (ay + 24)))
            hline(s, mx_ + mw + 32, ay + 24, 32)
            rect(s, ax, ay, aw, leaf_h, fill=PRIMARY_BG if a_focus else None,
                 line=LINE_ON_TINT if a_focus else LINE, radius=True)
            text(s, ax + 16, ay, aw - 96, leaf_h, act, size=SIZE_BODY, anchor=MSO_ANCHOR.MIDDLE)
            if a_focus:
                label(s, ax + aw - 80, ay + 8, "重点")
        y += block_h + 16
    return s
